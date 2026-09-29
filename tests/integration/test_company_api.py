import asyncio
import os
import re
from collections.abc import AsyncGenerator, AsyncIterator, Iterator, Mapping
from contextlib import asynccontextmanager
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from alembic import command
from alembic.config import Config
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import (
    AsyncConnection,
    AsyncEngine,
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)

from skmc_erp.config import Settings, get_settings


COMPANY_IMPORT_HEADERS = [
    "Legal Company Name",
    "Display / Short Name",
    "Base Time Zone",
    "Business Nature",
    "Company Email",
    "Company Phone",
    "Website",
]
FINANCIAL_YEAR_IMPORT_HEADERS = [
    "Fiscal Year Pattern",
    "Custom Start Month",
    "Custom Start Day",
    "Start Year",
]
LOCATION_IMPORT_HEADERS = [
    "Location Code",
    "Location Name",
    "Address Line 1",
    "Address Line 2",
    "City",
    "District",
    "Country Code",
    "Subdivision Code",
    "Postal Code",
    "Registered Office",
    "Corporate Office",
    "Branch",
    "Billing Office",
    "Warehouse",
    "Other Purpose",
]
GST_REGISTRATION_IMPORT_HEADERS = [
    "GSTIN",
    "Registered Legal Name",
    "Subdivision Code",
    "Registration Type Code",
    "Valid From",
    "Valid To",
]
GST_LOCATION_MAPPING_IMPORT_HEADERS = ["GSTIN", "Location Code"]


def _company_import_workbook(values: list[object]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Company Details"
    worksheet.append(COMPANY_IMPORT_HEADERS)
    worksheet.append(values)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _multi_sheet_import_workbook(
    *,
    company_values: list[object],
    financial_year_rows: list[list[object]],
    location_rows: list[list[object]],
    gst_registration_rows: list[list[object]] | None = None,
    gst_location_mapping_rows: list[list[object]] | None = None,
    order: tuple[str, ...] = ("Company Details", "Financial Years", "Locations"),
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows_by_sheet = {
        "Company Details": (COMPANY_IMPORT_HEADERS, [company_values]),
        "Financial Years": (FINANCIAL_YEAR_IMPORT_HEADERS, financial_year_rows),
        "Locations": (LOCATION_IMPORT_HEADERS, location_rows),
        "GST Registrations": (
            GST_REGISTRATION_IMPORT_HEADERS,
            gst_registration_rows or [],
        ),
        "GST Location Mappings": (
            GST_LOCATION_MAPPING_IMPORT_HEADERS,
            gst_location_mapping_rows or [],
        ),
    }
    for sheet_name in order:
        worksheet = workbook.create_sheet(sheet_name)
        headers, rows = rows_by_sheet[sheet_name]
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


@asynccontextmanager
async def _connection(database_url: str) -> AsyncIterator[AsyncConnection]:
    engine = create_async_engine(database_url)
    try:
        async with engine.connect() as connection:
            yield connection
    finally:
        await engine.dispose()


async def _assert_safe_empty_database(database_url: str) -> None:
    async with _connection(database_url) as connection:
        core_exists = await connection.scalar(
            text("SELECT to_regnamespace('core') IS NOT NULL")
        )
        if core_exists:
            pytest.fail("TEST_DATABASE_URL must point to a database without core schema")

        version_table_exists = await connection.scalar(
            text("SELECT to_regclass('public.alembic_version') IS NOT NULL")
        )
        if version_table_exists:
            version_count = await connection.scalar(
                text("SELECT count(*) FROM public.alembic_version")
            )
            if version_count:
                pytest.fail("TEST_DATABASE_URL already has an applied Alembic revision")


async def _assert_all_migrations_removed(database_url: str) -> None:
    async with _connection(database_url) as connection:
        core_exists = await connection.scalar(
            text("SELECT to_regnamespace('core') IS NOT NULL")
        )
        assert not core_exists


@pytest.fixture(scope="module")
def migrated_database() -> Iterator[str]:
    database_url = os.getenv("TEST_DATABASE_URL")
    if database_url is None:
        pytest.skip("TEST_DATABASE_URL is required for PostgreSQL API tests")

    parsed_url = make_url(database_url)
    if parsed_url.drivername != "postgresql+asyncpg":
        pytest.fail("TEST_DATABASE_URL must use postgresql+asyncpg")

    asyncio.run(_assert_safe_empty_database(database_url))

    previous_database_url = os.environ.get("DATABASE_URL")
    os.environ["DATABASE_URL"] = database_url
    get_settings.cache_clear()

    alembic_config = Config("alembic.ini")
    command.upgrade(alembic_config, "head")

    try:
        yield database_url
    finally:
        try:
            command.downgrade(alembic_config, "base")
            asyncio.run(_assert_all_migrations_removed(database_url))
        finally:
            if previous_database_url is None:
                os.environ.pop("DATABASE_URL", None)
            else:
                os.environ["DATABASE_URL"] = previous_database_url
            get_settings.cache_clear()


@pytest.fixture
async def api_context(
    migrated_database: str,
) -> AsyncGenerator[tuple[AsyncClient, AsyncEngine, FastAPI], None]:
    from skmc_erp.database import get_db_session
    from skmc_erp.main import app

    test_engine = create_async_engine(migrated_database)
    session_factory = async_sessionmaker(
        bind=test_engine,
        class_=AsyncSession,
    )

    async def override_db_session() -> AsyncGenerator[AsyncSession, None]:
        async with session_factory() as session:
            try:
                yield session
            except Exception:
                await session.rollback()
                raise

    development_settings = Settings(
        database_url=migrated_database,
        environment="development",
    )
    app.dependency_overrides[get_db_session] = override_db_session
    app.dependency_overrides[get_settings] = lambda: development_settings

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        try:
            yield client, test_engine, app
        finally:
            app.dependency_overrides.clear()
            async with test_engine.begin() as connection:
                await connection.execute(
                    text(
                        """
                        TRUNCATE TABLE
                            core.companies,
                            core.organisations,
                            core.country_subdivisions,
                            core.entity_types,
                            core.currencies,
                            core.countries,
                            core.tenants
                        CASCADE
                        """
                    )
                )
            await test_engine.dispose()


async def _execute(
    engine: AsyncEngine,
    sql: str,
    parameters: Mapping[str, object] | None = None,
) -> None:
    async with engine.begin() as connection:
        await connection.execute(text(sql), parameters or {})


async def _scalar(
    engine: AsyncEngine,
    sql: str,
    parameters: Mapping[str, object] | None = None,
) -> object:
    async with engine.connect() as connection:
        return await connection.scalar(text(sql), parameters or {})


async def _insert_tenant(
    engine: AsyncEngine,
    *,
    tenant_id: UUID,
    status: str = "ACTIVE",
) -> None:
    await _execute(
        engine,
        """
        INSERT INTO core.tenants (id, name, status)
        VALUES (:tenant_id, :name, :status)
        """,
        {
            "tenant_id": tenant_id,
            "name": f"Tenant {tenant_id}",
            "status": status,
        },
    )


async def _insert_organisation(
    engine: AsyncEngine,
    *,
    organisation_id: UUID,
    tenant_id: UUID,
    status: str = "ACTIVE",
) -> None:
    await _execute(
        engine,
        """
        INSERT INTO core.organisations (id, tenant_id, name, status)
        VALUES (:id, :tenant_id, :name, :status)
        """,
        {
            "id": organisation_id,
            "tenant_id": tenant_id,
            "name": f"Organisation {organisation_id}",
            "status": status,
        },
    )


async def _insert_country(
    engine: AsyncEngine,
    *,
    code: str,
    status: str = "ACTIVE",
) -> None:
    await _execute(
        engine,
        """
        INSERT INTO core.countries (code, name, status)
        VALUES (:code, :name, :status)
        """,
        {"code": code, "name": f"Country {code}", "status": status},
    )


async def _insert_subdivision(
    engine: AsyncEngine,
    *,
    country_code: str,
    code: str,
    status: str = "ACTIVE",
    gst_state_code: str | None = None,
) -> None:
    await _execute(
        engine,
        """
        INSERT INTO core.country_subdivisions
            (country_code, code, name, subdivision_type, gst_state_code, status)
        VALUES
            (:country_code, :code, :name, 'STATE', :gst_state_code, :status)
        """,
        {
            "country_code": country_code,
            "code": code,
            "name": f"Subdivision {code}",
            "gst_state_code": gst_state_code,
            "status": status,
        },
    )


async def _insert_entity_type(
    engine: AsyncEngine,
    *,
    entity_type_id: UUID,
    country_code: str,
    status: str = "ACTIVE",
) -> None:
    await _execute(
        engine,
        """
        INSERT INTO core.entity_types
            (id, country_code, code, name, status)
        VALUES
            (:id, :country_code, :code, :name, :status)
        """,
        {
            "id": entity_type_id,
            "country_code": country_code,
            "code": f"TYPE_{entity_type_id.hex.upper()}",
            "name": f"Entity Type {entity_type_id}",
            "status": status,
        },
    )


async def _insert_currency(
    engine: AsyncEngine,
    *,
    code: str,
    status: str = "ACTIVE",
) -> None:
    await _execute(
        engine,
        """
        INSERT INTO core.currencies
            (code, name, minor_units, status)
        VALUES
            (:code, :name, 2, :status)
        """,
        {"code": code, "name": f"Currency {code}", "status": status},
    )


async def test_tenant_header_validation_and_lookup(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, _, _ = api_context

    missing = await client.post("/companies", json={"legal_name": "Example"})
    malformed = await client.post(
        "/companies",
        headers={"X-Tenant-ID": "not-a-uuid"},
        json={"legal_name": "Example"},
    )
    nonexistent = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(uuid4())},
        json={"legal_name": "Example"},
    )

    assert missing.status_code == 422
    assert malformed.status_code == 422
    assert nonexistent.status_code == 404
    assert nonexistent.json() == {"detail": "Tenant not found"}


async def test_inactive_tenant_and_production_context_are_rejected(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, app = api_context
    inactive_tenant_id = uuid4()
    await _insert_tenant(
        engine,
        tenant_id=inactive_tenant_id,
        status="INACTIVE",
    )

    inactive = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(inactive_tenant_id)},
        json={"legal_name": "Example"},
    )
    assert inactive.status_code == 409
    assert inactive.json() == {"detail": "Tenant is not active"}

    production_settings = Settings(
        database_url="postgresql+asyncpg://example:example@localhost/example",
        environment="production",
    )
    app.dependency_overrides[get_settings] = lambda: production_settings
    production = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(inactive_tenant_id)},
        json={"legal_name": "Example"},
    )

    assert production.status_code == 503
    assert production.json() == {
        "detail": "Temporary tenant context is unavailable in production"
    }


async def test_minimal_draft_creation_persists_system_generated_values(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)

    first = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "  First Company  "},
    )
    second = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Second Company"},
    )

    assert first.status_code == 201
    assert second.status_code == 201
    first_body = first.json()
    second_body = second.json()
    assert UUID(first_body["id"])
    assert first_body["tenant_id"] == str(tenant_id)
    assert first_body["legal_name"] == "First Company"
    assert first_body["status"] == "DRAFT"
    assert re.fullmatch(r"COM\d{6,}", first_body["company_code"])
    assert first_body["company_code"] != second_body["company_code"]
    assert first_body["created_at"]
    assert first_body["updated_at"]
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.companies WHERE tenant_id = :tenant_id",
        {"tenant_id": tenant_id},
    ) == 2
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_legal_name_versions "
        "WHERE company_id = :company_id "
        "AND legal_name = 'First Company' "
        "AND valid_from = CURRENT_DATE "
        "AND valid_to IS NULL",
        {"company_id": UUID(first_body["id"])},
    ) == 1


async def test_active_optional_references_succeed(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    organisation_id = uuid4()
    entity_type_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_organisation(
        engine,
        organisation_id=organisation_id,
        tenant_id=tenant_id,
    )
    await _insert_country(engine, code="IN")
    await _insert_entity_type(
        engine,
        entity_type_id=entity_type_id,
        country_code="IN",
    )
    await _insert_currency(engine, code="INR")

    response = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "legal_name": "Complete Draft",
            "organisation_id": str(organisation_id),
            "entity_type_id": str(entity_type_id),
            "country_code": "IN",
            "base_currency_code": "INR",
            "base_timezone": "Asia/Kolkata",
            "business_nature": "BOTH",
        },
    )

    assert response.status_code == 201
    assert response.json()["organisation_id"] == str(organisation_id)
    assert response.json()["entity_type_id"] == str(entity_type_id)
    assert response.json()["country_code"] == "IN"
    assert response.json()["base_currency_code"] == "INR"
    assert response.json()["base_timezone"] == "Asia/Kolkata"


async def test_organisation_tenant_and_status_rules(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    cross_tenant_organisation_id = uuid4()
    inactive_organisation_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)
    await _insert_organisation(
        engine,
        organisation_id=cross_tenant_organisation_id,
        tenant_id=other_tenant_id,
    )
    await _insert_organisation(
        engine,
        organisation_id=inactive_organisation_id,
        tenant_id=tenant_id,
        status="INACTIVE",
    )

    cross_tenant = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "legal_name": "Cross Tenant",
            "organisation_id": str(cross_tenant_organisation_id),
        },
    )
    inactive = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "legal_name": "Inactive Organisation",
            "organisation_id": str(inactive_organisation_id),
        },
    )

    assert cross_tenant.status_code == 422
    assert inactive.status_code == 409
    assert await _scalar(engine, "SELECT count(*) FROM core.companies") == 0


async def test_missing_references_are_rejected(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)

    payloads = [
        {"organisation_id": str(uuid4())},
        {"entity_type_id": str(uuid4())},
        {"country_code": "ZZ"},
        {"base_currency_code": "ZZZ"},
    ]
    for index, reference in enumerate(payloads):
        response = await client.post(
            "/companies",
            headers={"X-Tenant-ID": str(tenant_id)},
            json={"legal_name": f"Missing Reference {index}", **reference},
        )
        assert response.status_code == 422

    assert await _scalar(engine, "SELECT count(*) FROM core.companies") == 0


async def test_inactive_references_and_country_mismatch_are_rejected(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    inactive_entity_type_id = uuid4()
    active_entity_type_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN", status="INACTIVE")
    await _insert_country(engine, code="US")
    await _insert_entity_type(
        engine,
        entity_type_id=inactive_entity_type_id,
        country_code="US",
        status="INACTIVE",
    )
    await _insert_entity_type(
        engine,
        entity_type_id=active_entity_type_id,
        country_code="US",
    )
    await _insert_currency(engine, code="INR", status="INACTIVE")

    payloads = [
        ({"entity_type_id": str(inactive_entity_type_id)}, 409),
        ({"country_code": "IN"}, 409),
        ({"base_currency_code": "INR"}, 409),
        (
            {
                "entity_type_id": str(active_entity_type_id),
                "country_code": "IN",
            },
            409,
        ),
    ]
    for index, (reference, expected_status) in enumerate(payloads):
        response = await client.post(
            "/companies",
            headers={"X-Tenant-ID": str(tenant_id)},
            json={"legal_name": f"Rejected Reference {index}", **reference},
        )
        assert response.status_code == expected_status

    await _execute(
        engine,
        "UPDATE core.countries SET status = 'ACTIVE' WHERE code = 'IN'",
    )
    mismatch = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "legal_name": "Country Mismatch",
            "entity_type_id": str(active_entity_type_id),
            "country_code": "IN",
        },
    )
    assert mismatch.status_code == 422
    assert await _scalar(engine, "SELECT count(*) FROM core.companies") == 0


async def test_company_reads_are_tenant_scoped(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)

    first = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "First Tenant Company"},
    )
    second = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Second Tenant Company"},
    )
    other = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(other_tenant_id)},
        json={"legal_name": "Other Tenant Company"},
    )
    assert first.status_code == second.status_code == other.status_code == 201

    listed = await client.get(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    fetched = await client.get(
        f"/companies/{first.json()['id']}",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    concealed = await client.get(
        f"/companies/{other.json()['id']}",
        headers={"X-Tenant-ID": str(tenant_id)},
    )

    assert listed.status_code == 200
    assert {company["id"] for company in listed.json()} == {
        first.json()["id"],
        second.json()["id"],
    }
    assert fetched.status_code == 200
    assert fetched.json()["legal_name"] == "First Tenant Company"
    assert concealed.status_code == 404
    assert concealed.json() == {"detail": "Company not found"}


async def test_company_profile_update_is_narrow_and_tenant_scoped(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Stable Legal Name"},
    )
    assert created.status_code == 201
    company_id = created.json()["id"]
    await _execute(
        engine,
        "UPDATE core.company_legal_name_versions "
        "SET valid_from = CURRENT_DATE - 1 "
        "WHERE company_id = :company_id",
        {"company_id": UUID(company_id)},
    )

    updated = await client.patch(
        f"/companies/{company_id}",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "legal_name": "  Updated Legal Name  ",
            "display_name": "  Updated Display  ",
            "email": "  finance@example.test  ",
            "phone": "  +91 120 555 0100  ",
            "website": "  https://example.test  ",
            "base_timezone": "Asia/Kolkata",
            "business_nature": "BOTH",
        },
    )
    concealed = await client.patch(
        f"/companies/{company_id}",
        headers={"X-Tenant-ID": str(other_tenant_id)},
        json={"display_name": "Hidden"},
    )
    history = await client.get(
        f"/companies/{company_id}/legal-name-history",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    concealed_history = await client.get(
        f"/companies/{company_id}/legal-name-history",
        headers={"X-Tenant-ID": str(other_tenant_id)},
    )
    same_day_change = await client.patch(
        f"/companies/{company_id}",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Second Same-Day Legal Name"},
    )

    assert updated.status_code == 200
    assert updated.json()["display_name"] == "Updated Display"
    assert updated.json()["email"] == "finance@example.test"
    assert updated.json()["base_timezone"] == "Asia/Kolkata"
    assert updated.json()["business_nature"] == "BOTH"
    assert updated.json()["legal_name"] == "Updated Legal Name"
    assert concealed.status_code == 404
    assert history.status_code == 200
    assert [row["legal_name"] for row in history.json()] == [
        "Stable Legal Name",
        "Updated Legal Name",
    ]
    assert history.json()[0]["valid_to"] < history.json()[1]["valid_from"]
    assert history.json()[1]["valid_to"] is None
    assert concealed_history.status_code == 404
    assert same_day_change.status_code == 409
    assert await _scalar(
        engine,
        "SELECT legal_name FROM core.companies WHERE id = :id",
        {"id": UUID(company_id)},
    ) == "Updated Legal Name"


@pytest.mark.parametrize(
    "field_name",
    [
        "id",
        "tenant_id",
        "company_code",
        "organisation_id",
        "entity_type_id",
        "country_code",
        "base_currency_code",
        "status",
        "created_at",
    ],
)
async def test_company_patch_rejects_unapproved_fields(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
    field_name: str,
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Immutable Company"},
    )

    response = await client.patch(
        f"/companies/{created.json()['id']}",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={field_name: "not-allowed"},
    )

    assert response.status_code == 422


async def test_company_patch_rejects_invalid_timezone(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Timezone Company"},
    )

    response = await client.patch(
        f"/companies/{created.json()['id']}",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"base_timezone": "Not/A-Timezone"},
    )

    assert response.status_code == 422


async def test_company_inactivation_is_controlled_and_terminal(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Lifecycle Company"},
    )
    company_id = created.json()["id"]

    draft_rejected = await client.post(
        f"/companies/{company_id}/inactivate",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    assert draft_rejected.status_code == 409

    await _execute(
        engine,
        "UPDATE core.companies SET status = 'ACTIVE' WHERE id = :id",
        {"id": UUID(company_id)},
    )
    inactivated = await client.post(
        f"/companies/{company_id}/inactivate",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    repeated = await client.post(
        f"/companies/{company_id}/inactivate",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    mutation = await client.patch(
        f"/companies/{company_id}",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Forbidden Rename"},
    )
    history = await client.get(
        f"/companies/{company_id}/legal-name-history",
        headers={"X-Tenant-ID": str(tenant_id)},
    )
    concealed = await client.post(
        f"/companies/{company_id}/inactivate",
        headers={"X-Tenant-ID": str(other_tenant_id)},
    )

    assert inactivated.status_code == 200
    assert inactivated.json()["status"] == "INACTIVE"
    assert repeated.status_code == 200
    assert mutation.status_code == 409
    assert history.status_code == 200
    assert [row["legal_name"] for row in history.json()] == [
        "Lifecycle Company"
    ]
    assert history.json()[0]["valid_to"] is None
    assert concealed.status_code == 404


async def test_company_import_preview_apply_and_idempotent_replay(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Import Company"},
    )
    company_id = created.json()["id"]
    contents = _company_import_workbook(
        [
            "Import Company",
            "Imported Display",
            "Asia/Kolkata",
            "BOTH",
            "finance@example.com",
            "+91 98912 55499",
            "https://example.com",
        ]
    )

    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "company-details.xlsx",
                contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert preview.status_code == 200
    preview_body = preview.json()
    assert preview_body["valid"] is True
    assert preview_body["normalized_values"]["phone"] == "+919891255499"
    assert set(preview_body["update_fields"]) == {
        "display_name",
        "base_timezone",
        "business_nature",
        "email",
        "phone",
        "website",
    }
    assert await _scalar(
        engine,
        "SELECT display_name FROM core.companies WHERE id = :id",
        {"id": UUID(company_id)},
    ) is None

    applied = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview_body["preview_token"]},
    )
    replayed = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview_body["preview_token"]},
    )
    second_preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "company-details.xlsx",
                contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert applied.status_code == 200
    assert applied.json()["no_changes"] is False
    assert applied.json()["company"]["display_name"] == "Imported Display"
    assert replayed.status_code == 200
    assert replayed.json()["no_changes"] is True
    assert second_preview.status_code == 200
    assert second_preview.json()["update_fields"] == []
    assert await _scalar(
        engine,
        "SELECT display_name FROM core.companies WHERE id = :id",
        {"id": UUID(company_id)},
    ) == "Imported Display"


async def test_company_import_rejects_cross_tenant_access_and_non_xlsx(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Private Import Company"},
    )
    company_id = created.json()["id"]

    concealed = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(other_tenant_id)},
        files={"workbook": ("payload.txt", b"not xlsx", "text/plain")},
    )
    unsupported = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("payload.txt", b"not xlsx", "text/plain")},
    )

    assert concealed.status_code == 404
    assert unsupported.status_code == 415
    assert "Only .xlsx" in unsupported.json()["detail"][0]["message"]


async def test_company_import_apply_uses_signed_server_authoritative_state(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Authoritative Import Company"},
    )
    company_id = created.json()["id"]
    contents = _company_import_workbook(
        [
            "Authoritative Import Company",
            "Signed Display",
            None,
            None,
            None,
            None,
            None,
        ]
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "company-details.xlsx",
                contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    token = preview.json()["preview_token"]

    extra_client_state = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "preview_token": token,
            "normalized_values": {"display_name": "Forged Display"},
        },
    )
    tampered = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": token + "x"},
    )

    assert extra_client_state.status_code == 422
    assert tampered.status_code == 422
    assert await _scalar(
        engine,
        "SELECT display_name FROM core.companies WHERE id = :id",
        {"id": UUID(company_id)},
    ) is None


async def test_company_import_apply_rolls_back_service_failure(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from skmc_erp.core.company.model import Company
    from skmc_erp.core.company.service import CompanyStateConflictError
    from skmc_erp.core.company_import import service as import_service

    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Rollback Import Company"},
    )
    company_id = created.json()["id"]
    contents = _company_import_workbook(
        [
            "Rollback Import Company",
            "Must Roll Back",
            None,
            None,
            None,
            None,
            None,
        ]
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "company-details.xlsx",
                contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    async def failing_update_company(**kwargs: object) -> Company:
        session = kwargs["session"]
        target_id = kwargs["company_id"]
        assert isinstance(session, AsyncSession)
        target = await session.get(Company, target_id)
        assert target is not None
        target.display_name = "Partially Written"
        await session.flush()
        raise CompanyStateConflictError("simulated apply failure")

    monkeypatch.setattr(import_service, "update_company", failing_update_company)
    failed = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview.json()["preview_token"]},
    )

    assert failed.status_code == 409
    assert await _scalar(
        engine,
        "SELECT display_name FROM core.companies WHERE id = :id",
        {"id": UUID(company_id)},
    ) is None


async def test_company_import_multi_sheet_preview_apply_and_replay(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(engine, country_code="IN", code="IN-UP")
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Multi Sheet Import Company"},
    )
    company_id = created.json()["id"]
    contents = _multi_sheet_import_workbook(
        company_values=[
            "Multi Sheet Import Company",
            "Configured Company",
            "Asia/Kolkata",
            "BOTH",
            "finance@example.com",
            "+919891255499",
            "https://example.com",
        ],
        financial_year_rows=[
            ["apr_mar", None, None, 2025],
            ["APR_MAR", None, None, 2026],
        ],
        location_rows=[
            [
                " noida_ro ",
                "Noida Office",
                "Sector 62",
                None,
                "Noida",
                None,
                "in",
                "in-up",
                "201309",
                "YES",
                "YES",
                "NO",
                "YES",
                "NO",
                None,
            ],
            [
                None,
                "Noida Warehouse",
                "Sector 63",
                None,
                "Noida",
                None,
                "IN",
                "IN-UP",
                "201301",
                "NO",
                "NO",
                "NO",
                "NO",
                "YES",
                None,
            ],
        ],
        order=("Locations", "Company Details", "Financial Years"),
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "configuration.xlsx",
                contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["valid"] is True
    assert body["sheets"] == ["Locations", "Company Details", "Financial Years"]
    assert body["fiscal_settings"]["classification"] == "NEW"
    assert [row["classification"] for row in body["financial_years"]] == [
        "NEW",
        "NEW",
    ]
    assert body["financial_years"][0]["display_code"] == "FY2025-26"
    assert [row["classification"] for row in body["locations"]] == [
        "NEW",
        "NEW",
    ]
    assert body["locations"][1]["code_will_be_generated"] is True

    applied = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": body["preview_token"]},
    )
    replayed = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": body["preview_token"]},
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["no_changes"] is False
    assert [row["location_code"] for row in applied.json()["locations"]] == [
        "NOIDA_RO",
        "LOC-0001",
    ]
    assert replayed.status_code == 200, replayed.text
    assert replayed.json()["no_changes"] is True
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.financial_years WHERE company_id = :company",
        {"company": UUID(company_id)},
    ) == 2
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_locations WHERE company_id = :company",
        {"company": UUID(company_id)},
    ) == 2


async def test_company_import_multi_sheet_failure_rolls_back_every_sheet(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from skmc_erp.core.company_import import service as import_service
    from skmc_erp.core.company_location.service import (
        CompanyLocationStateConflictError,
    )

    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Atomic Import Company", "country_code": "IN"},
    )
    company_id = created.json()["id"]
    contents = _multi_sheet_import_workbook(
        company_values=[
            "Atomic Import Company",
            "Must Roll Back",
            None,
            None,
            None,
            None,
            None,
        ],
        financial_year_rows=[["JAN_DEC", None, None, 2026]],
        location_rows=[
            [
                "ROLLBACK_BRANCH",
                "Rollback Branch",
                "Address",
                None,
                "Noida",
                None,
                "IN",
                "IN-UP",
                None,
                "NO",
                "NO",
                "YES",
                "NO",
                "NO",
                None,
            ]
        ],
        gst_registration_rows=[
            [
                "09ABCDE1234F1Z5",
                "Atomic Import Company",
                "IN-UP",
                None,
                None,
                None,
            ]
        ],
        gst_location_mapping_rows=[
            ["09ABCDE1234F1Z5", "ROLLBACK_BRANCH"]
        ],
        order=(
            "GST Location Mappings",
            "Locations",
            "Company Details",
            "GST Registrations",
            "Financial Years",
        ),
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("configuration.xlsx", contents, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview.status_code == 200, preview.text

    original_update = import_service.update_company_location

    async def fail_mapping_update(**kwargs: object) -> object:
        location_data = kwargs["location_data"]
        if getattr(location_data, "gst_registration_id", None) is not None:
            raise CompanyLocationStateConflictError("simulated final-row failure")
        return await original_update(**kwargs)

    monkeypatch.setattr(import_service, "update_company_location", fail_mapping_update)
    applied = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview.json()["preview_token"]},
    )
    assert applied.status_code == 409
    assert await _scalar(
        engine,
        "SELECT display_name FROM core.companies WHERE id = :id",
        {"id": UUID(company_id)},
    ) is None
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_fiscal_settings WHERE company_id = :id",
        {"id": UUID(company_id)},
    ) == 0
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.financial_years WHERE company_id = :id",
        {"id": UUID(company_id)},
    ) == 0
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_locations WHERE company_id = :id",
        {"id": UUID(company_id)},
    ) == 0
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_gst_registrations WHERE company_id = :id",
        {"id": UUID(company_id)},
    ) == 0


async def test_company_import_gst_create_mapping_and_replay(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    await _execute(
        engine,
        "INSERT INTO core.gst_registration_types (code, name, status) "
        "VALUES ('REGULAR', 'Regular', 'ACTIVE')",
    )
    created = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "GST Import Company", "country_code": "IN"},
    )
    company_id = created.json()["id"]
    workbook = _multi_sheet_import_workbook(
        company_values=[
            "GST Import Company",
            "GST Import",
            "Asia/Kolkata",
            "BOTH",
            None,
            None,
            None,
        ],
        financial_year_rows=[],
        location_rows=[
            [
                " del-ho ",
                "Delhi Head Office",
                "Connaught Place",
                None,
                "Delhi",
                None,
                "IN",
                "IN-UP",
                "110001",
                "NO",
                "YES",
                "NO",
                "YES",
                "NO",
                None,
            ]
        ],
        gst_registration_rows=[
            [
                " 09abcde1234f1z5 ",
                " GST Import Company Registered ",
                " in-up ",
                " regular ",
                "2026-04-01",
                "2027-03-31",
            ],
            [
                "09ABCDE1234F2Z4",
                None,
                "IN-UP",
                None,
                None,
                "2027-03-31",
            ],
        ],
        gst_location_mapping_rows=[
            ["09abcde1234f1z5", "del-ho"]
        ],
        order=(
            "GST Location Mappings",
            "GST Registrations",
            "Locations",
            "Company Details",
        ),
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "configuration.xlsx",
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert preview.status_code == 200, preview.text
    preview_body = preview.json()
    assert preview_body["valid"] is True
    assert preview_body["gst_registrations"][0]["gstin"] == "09ABCDE1234F1Z5"
    assert [
        row["classification"] for row in preview_body["gst_registrations"]
    ] == ["NEW", "NEW"]
    assert preview_body["gst_location_mappings"][0]["classification"] == "SAFE_UPDATE"

    applied = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview_body["preview_token"]},
    )
    replayed = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview_body["preview_token"]},
    )

    assert applied.status_code == 200, applied.text
    assert applied.json()["no_changes"] is False
    assert replayed.status_code == 200, replayed.text
    assert replayed.json()["no_changes"] is True
    assert replayed.json()["gst_registrations"][0]["classification"] == "UNCHANGED"
    assert replayed.json()["gst_location_mappings"][0]["classification"] == "UNCHANGED"
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_gst_registrations "
        "WHERE company_id = :company AND status = 'DRAFT'",
        {"company": UUID(company_id)},
    ) == 2
    assert await _scalar(
        engine,
        "SELECT registered_legal_name IS NULL AND gst_registration_type_id IS NULL "
        "AND valid_from IS NULL AND valid_to = DATE '2027-03-31' "
        "FROM core.company_gst_registrations WHERE gstin = :gstin",
        {"gstin": "09ABCDE1234F2Z4"},
    ) is True
    assert await _scalar(
        engine,
        "SELECT gst_registration_id IS NOT NULL FROM core.company_locations "
        "WHERE company_id = :company AND location_code = 'DEL-HO'",
        {"company": UUID(company_id)},
    ) is True


async def test_company_import_existing_gst_is_compare_only_and_concealed(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    selected = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Selected GST Company", "country_code": "IN"},
    )
    owner = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(other_tenant_id)},
        json={"legal_name": "Hidden GST Owner", "country_code": "IN"},
    )
    selected_id = selected.json()["id"]
    owner_id = owner.json()["id"]
    await _execute(
        engine,
        "INSERT INTO core.company_gst_registrations "
        "(company_id, gstin, registered_legal_name, subdivision_code, status) "
        "VALUES (:company, '09ABCDE1234F1Z5', 'Selected Registered', 'IN-UP', 'INACTIVE'), "
        "(:owner, '09ABCDE1234F2Z4', 'Actual Hidden Registered', 'IN-UP', 'DRAFT')",
        {"company": UUID(selected_id), "owner": UUID(owner_id)},
    )

    def gst_workbook(gstin: str, legal_name: str) -> bytes:
        return _multi_sheet_import_workbook(
            company_values=[],
            financial_year_rows=[],
            location_rows=[],
            gst_registration_rows=[
                [gstin, legal_name, "IN-UP", None, None, None]
            ],
            order=("GST Registrations",),
        )

    exact = await client.post(
        f"/companies/{selected_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("gst.xlsx", gst_workbook("09ABCDE1234F1Z5", "Selected Registered"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    changed = await client.post(
        f"/companies/{selected_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("gst.xlsx", gst_workbook("09ABCDE1234F1Z5", "Changed Name"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    foreign = await client.post(
        f"/companies/{selected_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("gst.xlsx", gst_workbook("09ABCDE1234F2Z4", "Workbook Supplied Name"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert exact.status_code == 200
    assert exact.json()["valid"] is True
    assert exact.json()["gst_registrations"][0]["classification"] == "UNCHANGED"
    assert changed.status_code == 200
    assert changed.json()["valid"] is False
    assert changed.json()["preview_token"] is None
    assert changed.json()["gst_registrations"][0]["classification"] == "CONFLICT"
    assert foreign.status_code == 200
    foreign_body = foreign.json()
    assert foreign_body["valid"] is False
    serialized = foreign.text
    assert owner_id not in serialized
    assert str(other_tenant_id) not in serialized
    assert "Hidden GST Owner" not in serialized
    assert "Actual Hidden Registered" not in serialized


async def test_company_import_gst_validation_and_mapping_conflicts(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    await _execute(
        engine,
        "INSERT INTO core.gst_registration_types (code, name, status) "
        "VALUES ('RETIRED', 'Retired', 'INACTIVE')",
    )
    company = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "GST Validation Company", "country_code": "IN"},
    )
    company_id = company.json()["id"]
    headers = {"X-Tenant-ID": str(tenant_id)}

    invalid_rows = [
        [
            ["09ABCDE1234F1Z5", None, "IN-UP", None, None, None],
            [" 09abcde1234f1z5 ", None, "IN-UP", None, None, None],
        ],
        [["09ABCDE1234F1Z5", None, "IN-UP", "MISSING", None, None]],
        [["09ABCDE1234F1Z5", None, "IN-UP", "RETIRED", None, None]],
        [["27ABCDE1234F1Z5", None, "IN-UP", None, None, None]],
        [["09ABCDE1234F1Z5", None, "IN-UP", None, "2027-01-01", "2026-01-01"]],
    ]
    for rows in invalid_rows:
        workbook = _multi_sheet_import_workbook(
            company_values=[],
            financial_year_rows=[],
            location_rows=[],
            gst_registration_rows=rows,
            order=("GST Registrations",),
        )
        response = await client.post(
            f"/companies/{company_id}/configuration-import/preview",
            headers=headers,
            files={"workbook": ("gst.xlsx", workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 422, response.text

    duplicate_mapping = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[],
        gst_location_mapping_rows=[
            ["09ABCDE1234F1Z5", "DEL-HO"],
            [" 09abcde1234f1z5 ", " del-ho "],
        ],
        order=("GST Location Mappings",),
    )
    duplicate_response = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers=headers,
        files={"workbook": ("mapping.xlsx", duplicate_mapping, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert duplicate_response.status_code == 422
    assert "Duplicate GST Location mapping" in duplicate_response.text


async def test_company_import_stale_location_mapping_is_not_overwritten(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    company = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Stale GST Company", "country_code": "IN"},
    )
    company_id = company.json()["id"]
    await _execute(
        engine,
        "INSERT INTO core.company_gst_registrations "
        "(company_id, gstin, subdivision_code, status) VALUES "
        "(:company, '09ABCDE1234F1Z5', 'IN-UP', 'DRAFT'), "
        "(:company, '09ABCDE1234F2Z4', 'IN-UP', 'INACTIVE')",
        {"company": UUID(company_id)},
    )
    location = await client.post(
        f"/companies/{company_id}/locations",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={
            "location_code": "STALE_LOC",
            "location_name": "Stale Location",
            "address_line_1": "Address",
            "city": "Noida",
            "country_code": "IN",
            "subdivision_code": "IN-UP",
            "is_branch": True,
        },
    )
    competing_workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[],
        gst_location_mapping_rows=[
            ["09ABCDE1234F1Z5", "STALE_LOC"],
            ["09ABCDE1234F2Z4", "STALE_LOC"],
        ],
        order=("GST Location Mappings",),
    )
    competing = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("mapping.xlsx", competing_workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert competing.status_code == 200
    assert competing.json()["valid"] is False
    assert {
        row["classification"]
        for row in competing.json()["gst_location_mappings"]
    } == {"CONFLICT"}
    inactive_workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[],
        gst_location_mapping_rows=[["09ABCDE1234F2Z4", "STALE_LOC"]],
        order=("GST Location Mappings",),
    )
    inactive_preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("mapping.xlsx", inactive_workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert inactive_preview.status_code == 200
    assert inactive_preview.json()["valid"] is False
    assert "Inactive GST Registration" in inactive_preview.text
    workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[],
        gst_location_mapping_rows=[["09ABCDE1234F1Z5", "STALE_LOC"]],
        order=("GST Location Mappings",),
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={"workbook": ("mapping.xlsx", workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview.status_code == 200
    assert preview.json()["gst_location_mappings"][0]["classification"] == "SAFE_UPDATE"
    other_id = await _scalar(
        engine,
        "SELECT id FROM core.company_gst_registrations WHERE gstin = '09ABCDE1234F2Z4'",
    )
    await _execute(
        engine,
        "UPDATE core.company_locations SET gst_registration_id = :gst WHERE id = :location",
        {"gst": other_id, "location": UUID(location.json()["id"])},
    )

    applied = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview.json()["preview_token"]},
    )

    assert applied.status_code == 409
    assert await _scalar(
        engine,
        "SELECT gst_registration_id FROM core.company_locations WHERE id = :id",
        {"id": UUID(location.json()["id"])},
    ) == other_id


async def test_company_import_rejects_gstin_created_elsewhere_after_preview(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    other_tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_tenant(engine, tenant_id=other_tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    selected = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "Selected Stale Company", "country_code": "IN"},
    )
    foreign = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(other_tenant_id)},
        json={"legal_name": "Concealed Concurrent Owner", "country_code": "IN"},
    )
    workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[],
        gst_registration_rows=[
            ["09ABCDE1234F1Z5", None, "IN-UP", None, None, None]
        ],
        order=("GST Registrations",),
    )
    preview = await client.post(
        f"/companies/{selected.json()['id']}/configuration-import/preview",
        headers={"X-Tenant-ID": str(tenant_id)},
        files={
            "workbook": (
                "gst.xlsx",
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200
    assert preview.json()["gst_registrations"][0]["classification"] == "NEW"

    await _execute(
        engine,
        "INSERT INTO core.company_gst_registrations "
        "(company_id, gstin, registered_legal_name, subdivision_code, status) "
        "VALUES (:company, '09ABCDE1234F1Z5', 'Concealed Concurrent GST', "
        "'IN-UP', 'DRAFT')",
        {"company": UUID(foreign.json()["id"])},
    )
    applied = await client.post(
        f"/companies/{selected.json()['id']}/configuration-import/apply",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"preview_token": preview.json()["preview_token"]},
    )

    assert applied.status_code == 409
    assert foreign.json()["id"] not in applied.text
    assert str(other_tenant_id) not in applied.text
    assert "Concealed Concurrent" not in applied.text
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_gst_registrations "
        "WHERE company_id = :company",
        {"company": UUID(selected.json()["id"])},
    ) == 0


async def test_company_import_gst_same_workbook_reference_combinations(
    api_context: tuple[AsyncClient, AsyncEngine, FastAPI],
) -> None:
    client, engine, _ = api_context
    tenant_id = uuid4()
    await _insert_tenant(engine, tenant_id=tenant_id)
    await _insert_country(engine, code="IN")
    await _insert_subdivision(
        engine,
        country_code="IN",
        code="IN-UP",
        gst_state_code="09",
    )
    company = await client.post(
        "/companies",
        headers={"X-Tenant-ID": str(tenant_id)},
        json={"legal_name": "GST Dependency Company", "country_code": "IN"},
    )
    company_id = company.json()["id"]
    headers = {"X-Tenant-ID": str(tenant_id)}
    await _execute(
        engine,
        "INSERT INTO core.company_gst_registrations "
        "(company_id, gstin, subdivision_code, status) "
        "VALUES (:company, '09ABCDE1234F1Z5', 'IN-UP', 'DRAFT')",
        {"company": UUID(company_id)},
    )
    existing_location = await client.post(
        f"/companies/{company_id}/locations",
        headers=headers,
        json={
            "location_code": "EXIST_LOC",
            "location_name": "Existing Location",
            "address_line_1": "Address",
            "city": "Noida",
            "country_code": "IN",
            "subdivision_code": "IN-UP",
            "is_branch": True,
        },
    )
    assert existing_location.status_code == 201

    new_gst_workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[],
        gst_registration_rows=[
            ["09ABCDE1234F2Z4", None, "IN-UP", None, None, None]
        ],
        gst_location_mapping_rows=[
            ["09ABCDE1234F2Z4", "EXIST_LOC"]
        ],
        order=("GST Location Mappings", "GST Registrations"),
    )
    preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers=headers,
        files={"workbook": ("gst.xlsx", new_gst_workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    applied = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers=headers,
        json={"preview_token": preview.json()["preview_token"]},
    )
    assert applied.status_code == 200, applied.text

    new_location_workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[
            [
                "NEW_LOC",
                "New Location",
                "New Address",
                None,
                "Noida",
                None,
                "IN",
                "IN-UP",
                None,
                "NO",
                "NO",
                "YES",
                "NO",
                "NO",
                None,
            ]
        ],
        gst_location_mapping_rows=[
            ["09ABCDE1234F1Z5", "NEW_LOC"]
        ],
        order=("GST Location Mappings", "Locations"),
    )
    second_preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers=headers,
        files={"workbook": ("mapping.xlsx", new_location_workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    second_apply = await client.post(
        f"/companies/{company_id}/configuration-import/apply",
        headers=headers,
        json={"preview_token": second_preview.json()["preview_token"]},
    )
    assert second_apply.status_code == 200, second_apply.text
    assert await _scalar(
        engine,
        "SELECT count(*) FROM core.company_locations "
        "WHERE company_id = :company AND gst_registration_id IS NOT NULL",
        {"company": UUID(company_id)},
    ) == 2

    blank_code_workbook = _multi_sheet_import_workbook(
        company_values=[],
        financial_year_rows=[],
        location_rows=[
            [
                None,
                "Generated Code Location",
                "Address",
                None,
                "Noida",
                None,
                "IN",
                "IN-UP",
                None,
                "NO",
                "NO",
                "YES",
                "NO",
                "NO",
                None,
            ]
        ],
        gst_location_mapping_rows=[
            ["09ABCDE1234F1Z5", "LOC-9999"]
        ],
        order=("Locations", "GST Location Mappings"),
    )
    blank_preview = await client.post(
        f"/companies/{company_id}/configuration-import/preview",
        headers=headers,
        files={"workbook": ("mapping.xlsx", blank_code_workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert blank_preview.status_code == 200
    assert blank_preview.json()["valid"] is False
    assert "explicitly supplied code" in blank_preview.text
