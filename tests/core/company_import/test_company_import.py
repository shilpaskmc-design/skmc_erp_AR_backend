from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import Workbook

from skmc_erp.core.company.model import Company, CompanyBusinessNature, CompanyStatus
from skmc_erp.core.company_import.service import (
    CompanyImportValidationError,
    _normalize_financial_years,
    _normalize_gst_location_mappings,
    _normalize_gst_registrations,
    _yes_no,
    preview_company_import,
)
from skmc_erp.core.company_import.token import (
    PreviewTokenError,
    create_preview_token,
    read_preview_token,
)
from skmc_erp.core.company_import.workbook import (
    COLUMN_TO_FIELD,
    FINANCIAL_YEAR_COLUMN_TO_FIELD,
    GST_LOCATION_MAPPING_COLUMN_TO_FIELD,
    GST_REGISTRATION_COLUMN_TO_FIELD,
    LOCATION_COLUMN_TO_FIELD,
    ParsedWorkbookRow,
    CompanyWorkbookError,
    parse_company_details_workbook,
    parse_company_workbook,
)
from skmc_erp.core.geography.model import CountrySubdivisionStatus


HEADERS = list(COLUMN_TO_FIELD)
SIGNING_KEY = "test-company-import-signing-key-1234567890"


def workbook_bytes(
    values: list[object] | None = None,
    *,
    headers: list[str] | None = None,
    title: str = "Company Details",
    extra_rows: list[list[object]] | None = None,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    worksheet.append(headers or HEADERS)
    if values is not None:
        worksheet.append(values)
    for row in extra_rows or []:
        worksheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def multi_sheet_workbook_bytes() -> bytes:
    workbook = Workbook()
    company_sheet = workbook.active
    company_sheet.title = "Company Details"
    company_sheet.append(HEADERS)
    company_sheet.append(valid_values())
    location_sheet = workbook.create_sheet("Locations", 0)
    location_headers = list(reversed(LOCATION_COLUMN_TO_FIELD))
    location_sheet.append(location_headers)
    values = {
        "Location Code": "BRANCH_1",
        "Location Name": "Branch",
        "Address Line 1": "Address",
        "Address Line 2": None,
        "City": "Noida",
        "District": None,
        "Country Code": "IN",
        "Subdivision Code": None,
        "Postal Code": None,
        "Registered Office": "NO",
        "Corporate Office": "NO",
        "Branch": "YES",
        "Billing Office": "NO",
        "Warehouse": "NO",
        "Other Purpose": None,
    }
    location_sheet.append([values[header] for header in location_headers])
    financial_sheet = workbook.create_sheet("Financial Years", 1)
    financial_headers = list(reversed(FINANCIAL_YEAR_COLUMN_TO_FIELD))
    financial_sheet.append(financial_headers)
    financial_values = {
        "Fiscal Year Pattern": "APR_MAR",
        "Custom Start Month": None,
        "Custom Start Day": None,
        "Start Year": 2026,
    }
    financial_sheet.append(
        [financial_values[header] for header in financial_headers]
    )
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def company() -> Company:
    now = datetime.now(UTC)
    return Company(
        id=uuid4(),
        tenant_id=uuid4(),
        organisation_id=None,
        legal_name="Current Legal Name",
        display_name="Current Display",
        company_code="COM000001",
        entity_type_id=None,
        country_code="IN",
        email="current@example.com",
        phone="+919891255499",
        website="https://current.example.com",
        base_timezone="Asia/Kolkata",
        base_currency_code="INR",
        business_nature=CompanyBusinessNature.SERVICES,
        status=CompanyStatus.DRAFT,
        created_at=now,
        updated_at=now,
    )


def valid_values() -> list[object]:
    return [
        "  Updated Legal Name  ",
        "  Updated Display  ",
        "Asia/Kolkata",
        "BOTH",
        "  finance@example.com  ",
        "9891255499",
        "  https://example.com/company  ",
    ]


def test_parser_accepts_exact_company_details_contract() -> None:
    parsed = parse_company_details_workbook(workbook_bytes(valid_values()))

    assert parsed.values == {
        "legal_name": "Updated Legal Name",
        "display_name": "Updated Display",
        "base_timezone": "Asia/Kolkata",
        "business_nature": "BOTH",
        "email": "finance@example.com",
        "phone": "9891255499",
        "website": "https://example.com/company",
    }


def test_parser_accepts_known_sheets_in_any_order_and_flexible_headers() -> None:
    parsed = parse_company_workbook(multi_sheet_workbook_bytes())

    assert parsed.sheets == ["Locations", "Financial Years", "Company Details"]
    assert parsed.financial_years[0].values["start_year"] == 2026
    assert parsed.locations[0].values["location_code"] == "BRANCH_1"


def test_parser_accepts_gst_sheets_with_flexible_headers() -> None:
    workbook = Workbook()
    registration_sheet = workbook.active
    registration_sheet.title = "GST Registrations"
    registration_headers = list(reversed(GST_REGISTRATION_COLUMN_TO_FIELD))
    registration_sheet.append(registration_headers)
    registration_values = {
        "GSTIN": " 09abcde1234f1z5 ",
        "Registered Legal Name": "Example Private Limited",
        "Subdivision Code": "IN-UP",
        "Registration Type Code": None,
        "Valid From": "2026-04-01",
        "Valid To": None,
    }
    registration_sheet.append(
        [registration_values[header] for header in registration_headers]
    )
    mapping_sheet = workbook.create_sheet("GST Location Mappings", 0)
    mapping_headers = list(reversed(GST_LOCATION_MAPPING_COLUMN_TO_FIELD))
    mapping_sheet.append(mapping_headers)
    mapping_values = {"GSTIN": "09ABCDE1234F1Z5", "Location Code": "DEL-HO"}
    mapping_sheet.append([mapping_values[header] for header in mapping_headers])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    parsed = parse_company_workbook(stream.getvalue())

    assert parsed.sheets == ["GST Location Mappings", "GST Registrations"]
    assert parsed.gst_registrations[0].values["gstin"] == "09abcde1234f1z5"
    assert parsed.gst_location_mappings[0].values["location_code"] == "DEL-HO"


@pytest.mark.parametrize(
    "headers",
    [
        [header for header in GST_REGISTRATION_COLUMN_TO_FIELD if header != "GSTIN"],
        [*GST_REGISTRATION_COLUMN_TO_FIELD, "Status"],
        [*GST_REGISTRATION_COLUMN_TO_FIELD, "GSTIN"],
    ],
)
def test_gst_registration_headers_are_exact(headers: list[str]) -> None:
    with pytest.raises(CompanyWorkbookError):
        parse_company_workbook(
            workbook_bytes(
                [None] * len(headers),
                headers=headers,
                title="GST Registrations",
            )
        )


@pytest.mark.parametrize(
    "headers",
    [
        ["GSTIN"],
        ["GSTIN", "Location Code", "Status"],
        ["GSTIN", "Location Code", "GSTIN"],
    ],
)
def test_gst_location_mapping_headers_are_exact(headers: list[str]) -> None:
    with pytest.raises(CompanyWorkbookError):
        parse_company_workbook(
            workbook_bytes(
                [None] * len(headers),
                headers=headers,
                title="GST Location Mappings",
            )
        )


@pytest.mark.asyncio
async def test_gst_normalization_rejects_duplicate_normalized_gstin() -> None:
    class FakeSession:
        async def scalar(self, _: object) -> object:
            return SimpleNamespace(
                status=CountrySubdivisionStatus.ACTIVE,
                country_code="IN",
                gst_state_code="09",
            )

    values = {
        "gstin": "09ABCDE1234F1Z5",
        "registered_legal_name": None,
        "subdivision_code": "IN-UP",
        "registration_type_code": None,
        "valid_from": "2026-04-01",
        "valid_to": None,
    }
    normalized, issues = await _normalize_gst_registrations(
        FakeSession(),  # type: ignore[arg-type]
        company(),
        [
            ParsedWorkbookRow(row=2, values=values),
            ParsedWorkbookRow(
                row=3,
                values={
                    **values,
                    "gstin": " 09abcde1234f1z5 ",
                    "registration_type_code": "DIFFERENT",
                },
            ),
        ],
    )

    assert len(normalized) == 1
    assert normalized[0].data.gstin == "09ABCDE1234F1Z5"
    assert normalized[0].data.valid_from.isoformat() == "2026-04-01"
    assert "first seen at row 2" in issues[0].message


@pytest.mark.asyncio
async def test_gst_normalization_rejects_empty_required_gstin() -> None:
    class FakeSession:
        async def scalar(self, _: object) -> object:
            raise AssertionError("invalid GSTIN must fail before database lookup")

    normalized, issues = await _normalize_gst_registrations(
        FakeSession(),  # type: ignore[arg-type]
        company(),
        [
            ParsedWorkbookRow(
                row=2,
                values={
                    "gstin": None,
                    "registered_legal_name": None,
                    "subdivision_code": "IN-UP",
                    "registration_type_code": None,
                    "valid_from": None,
                    "valid_to": None,
                },
            )
        ],
    )

    assert normalized == []
    assert issues[0].field == "gstin"


def test_gst_mapping_normalization_rejects_duplicate_pair() -> None:
    normalized, issues = _normalize_gst_location_mappings(
        [
            ParsedWorkbookRow(
                row=2,
                values={"gstin": "09ABCDE1234F1Z5", "location_code": "DEL-HO"},
            ),
            ParsedWorkbookRow(
                row=3,
                values={
                    "gstin": " 09abcde1234f1z5 ",
                    "location_code": " del-ho ",
                },
            ),
        ]
    )

    assert len(normalized) == 1
    assert normalized[0].location_code == "DEL-HO"
    assert "Duplicate GST Location mapping" in issues[0].message


def test_financial_year_normalization_rejects_duplicates_and_feb_29() -> None:
    settings, rows, issues = _normalize_financial_years(
        [
            ParsedWorkbookRow(
                row=2,
                values={
                    "fiscal_year_pattern": "custom",
                    "custom_fiscal_year_start_month": 2,
                    "custom_fiscal_year_start_day": 29,
                    "start_year": 2026,
                },
            ),
            ParsedWorkbookRow(
                row=3,
                values={
                    "fiscal_year_pattern": "APR_MAR",
                    "custom_fiscal_year_start_month": None,
                    "custom_fiscal_year_start_day": None,
                    "start_year": 2026,
                },
            ),
            ParsedWorkbookRow(
                row=4,
                values={
                    "fiscal_year_pattern": "APR_MAR",
                    "custom_fiscal_year_start_month": None,
                    "custom_fiscal_year_start_day": None,
                    "start_year": 2026,
                },
            ),
        ]
    )

    assert settings is not None
    assert [row.start_year for row in rows] == [2026]
    assert any("valid in every calendar year" in issue.message for issue in issues)
    assert any("Duplicate Start Year" in issue.message for issue in issues)


def test_location_boolean_contract_accepts_only_yes_no() -> None:
    assert _yes_no(" yes ", field="Branch") is True
    assert _yes_no("NO", field="Branch") is False
    with pytest.raises(ValueError, match="YES or NO"):
        _yes_no("TRUE", field="Branch")


@pytest.mark.asyncio
async def test_preview_normalizes_values_and_does_not_mutate_company() -> None:
    existing = company()
    tenant = SimpleNamespace(id=existing.tenant_id)

    preview = await preview_company_import(
        company=existing,
        tenant=tenant,
        contents=workbook_bytes(valid_values()),
        signing_key=SIGNING_KEY,
    )

    assert preview.valid
    assert preview.normalized_values["phone"] == "+919891255499"
    assert preview.normalized_values["email"] == "finance@example.com"
    assert preview.normalized_values["business_nature"] == "BOTH"
    assert set(preview.update_fields) == {
        "legal_name",
        "display_name",
        "business_nature",
        "email",
        "website",
    }
    assert set(preview.unchanged_fields) == {"base_timezone", "phone"}
    assert {gap.field for gap in preview.unsupported_fields} == {
        "country",
        "entity_type",
        "base_currency",
        "msme_details",
        "legal_identifiers",
    }
    assert existing.legal_name == "Current Legal Name"
    assert existing.display_name == "Current Display"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("column", "bad_value", "field"),
    [
        ("Company Email", "invalid", "email"),
        ("Company Phone", "12345", "phone"),
        ("Website", "ftp://example.com", "website"),
        ("Base Time Zone", "Not/A-Timezone", "base_timezone"),
        ("Business Nature", "CONSULTING", "business_nature"),
    ],
)
async def test_preview_reports_field_validation_errors(
    column: str,
    bad_value: str,
    field: str,
) -> None:
    values = valid_values()
    values[HEADERS.index(column)] = bad_value
    existing = company()

    with pytest.raises(CompanyImportValidationError) as caught:
        await preview_company_import(
            company=existing,
            tenant=SimpleNamespace(id=existing.tenant_id),
            contents=workbook_bytes(values),
            signing_key=SIGNING_KEY,
        )

    issue = caught.value.issues[0]
    assert issue.sheet == "Company Details"
    assert issue.row == 2
    assert issue.column == column
    assert issue.field == field


@pytest.mark.asyncio
async def test_phone_validation_uses_current_company_country() -> None:
    existing = company()
    existing.country_code = "US"
    values = valid_values()
    values[HEADERS.index("Company Phone")] = "9891255499"

    with pytest.raises(CompanyImportValidationError) as caught:
        await preview_company_import(
            company=existing,
            tenant=SimpleNamespace(id=existing.tenant_id),
            contents=workbook_bytes(values),
            signing_key=SIGNING_KEY,
        )

    assert caught.value.issues[0].field == "phone"
    assert "country context" in caught.value.issues[0].message


@pytest.mark.parametrize(
    "contents",
    [
        workbook_bytes(None),
        workbook_bytes(
            [None, "Display", "Asia/Kolkata", "SERVICES", None, None, None]
        ),
    ],
)
def test_missing_company_details_record_or_legal_name_is_rejected(
    contents: bytes,
) -> None:
    with pytest.raises(CompanyWorkbookError) as caught:
        parse_company_details_workbook(contents)

    assert caught.value.issues[0].row == 2


def test_unknown_and_injection_columns_are_rejected() -> None:
    headers = [*HEADERS, "tenant_id", "company_id", "status"]
    with pytest.raises(CompanyWorkbookError) as caught:
        parse_company_details_workbook(
            workbook_bytes([*valid_values(), str(uuid4()), str(uuid4()), "ACTIVE"], headers=headers)
        )

    messages = {issue.message for issue in caught.value.issues}
    assert messages == {
        "Unsupported column 'tenant_id'",
        "Unsupported column 'company_id'",
        "Unsupported column 'status'",
    }


def test_multiple_company_records_are_rejected() -> None:
    with pytest.raises(CompanyWorkbookError) as caught:
        parse_company_details_workbook(
            workbook_bytes(valid_values(), extra_rows=[valid_values()])
        )

    assert "exactly one logical record" in caught.value.issues[0].message
    assert caught.value.issues[0].row == 3


@pytest.mark.parametrize(
    "contents",
    [
        workbook_bytes(valid_values(), title="Companies"),
        b"not an xlsx file",
    ],
)
def test_wrong_sheet_and_malformed_workbook_are_rejected(contents: bytes) -> None:
    with pytest.raises(CompanyWorkbookError):
        parse_company_details_workbook(contents)


def test_formula_cells_are_rejected() -> None:
    values = valid_values()
    values[HEADERS.index("Display / Short Name")] = "=1+1"

    with pytest.raises(CompanyWorkbookError) as caught:
        parse_company_details_workbook(workbook_bytes(values))

    assert "Formulas are not allowed" in caught.value.issues[0].message


def test_preview_token_rejects_tampering_and_expiry() -> None:
    token, _ = create_preview_token(
        signing_key=SIGNING_KEY,
        tenant_id=str(uuid4()),
        company_id=str(uuid4()),
        values={"legal_name": "Example"},
        company_fingerprint="fingerprint",
        now=100,
    )
    payload = read_preview_token(token, signing_key=SIGNING_KEY, now=101)
    assert payload.values == {"legal_name": "Example"}

    with pytest.raises(PreviewTokenError):
        read_preview_token(token + "x", signing_key=SIGNING_KEY, now=101)
    with pytest.raises(PreviewTokenError, match="expired"):
        read_preview_token(token, signing_key=SIGNING_KEY, now=1000 + 100)
