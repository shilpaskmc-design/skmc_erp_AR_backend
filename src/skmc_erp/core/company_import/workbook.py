from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from skmc_erp.core.company_import.schema import CompanyImportIssue


SHEET_NAME = "Company Details"
FINANCIAL_YEARS_SHEET = "Financial Years"
LOCATIONS_SHEET = "Locations"
GST_REGISTRATIONS_SHEET = "GST Registrations"
GST_LOCATION_MAPPINGS_SHEET = "GST Location Mappings"
SUPPORTED_SHEETS = {
    SHEET_NAME,
    FINANCIAL_YEARS_SHEET,
    LOCATIONS_SHEET,
    GST_REGISTRATIONS_SHEET,
    GST_LOCATION_MAPPINGS_SHEET,
}

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ARCHIVE_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_ENTRY_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 200
MAX_WORKSHEET_ROWS = 100
MAX_WORKSHEET_COLUMNS = 20
MAX_CELL_CHARACTERS = 2048

COLUMN_TO_FIELD = {
    "Legal Company Name": "legal_name",
    "Display / Short Name": "display_name",
    "Base Time Zone": "base_timezone",
    "Business Nature": "business_nature",
    "Company Email": "email",
    "Company Phone": "phone",
    "Website": "website",
}
FIELD_TO_COLUMN = {field: column for column, field in COLUMN_TO_FIELD.items()}

FINANCIAL_YEAR_COLUMN_TO_FIELD = {
    "Fiscal Year Pattern": "fiscal_year_pattern",
    "Custom Start Month": "custom_fiscal_year_start_month",
    "Custom Start Day": "custom_fiscal_year_start_day",
    "Start Year": "start_year",
}
FINANCIAL_YEAR_FIELD_TO_COLUMN = {
    field: column for column, field in FINANCIAL_YEAR_COLUMN_TO_FIELD.items()
}

LOCATION_COLUMN_TO_FIELD = {
    "Location Code": "location_code",
    "Location Name": "location_name",
    "Address Line 1": "address_line_1",
    "Address Line 2": "address_line_2",
    "City": "city",
    "District": "district",
    "Country Code": "country_code",
    "Subdivision Code": "subdivision_code",
    "Postal Code": "postal_code",
    "Registered Office": "is_registered_office",
    "Corporate Office": "is_corporate_office",
    "Branch": "is_branch",
    "Billing Office": "is_billing_office",
    "Warehouse": "is_warehouse",
    "Other Purpose": "other_purpose",
}
LOCATION_FIELD_TO_COLUMN = {
    field: column for column, field in LOCATION_COLUMN_TO_FIELD.items()
}

GST_REGISTRATION_COLUMN_TO_FIELD = {
    "GSTIN": "gstin",
    "Registered Legal Name": "registered_legal_name",
    "Subdivision Code": "subdivision_code",
    "Registration Type Code": "registration_type_code",
    "Valid From": "valid_from",
    "Valid To": "valid_to",
}
GST_REGISTRATION_FIELD_TO_COLUMN = {
    field: column for column, field in GST_REGISTRATION_COLUMN_TO_FIELD.items()
}

GST_LOCATION_MAPPING_COLUMN_TO_FIELD = {
    "GSTIN": "gstin",
    "Location Code": "location_code",
}
GST_LOCATION_MAPPING_FIELD_TO_COLUMN = {
    field: column for column, field in GST_LOCATION_MAPPING_COLUMN_TO_FIELD.items()
}


@dataclass(frozen=True)
class ParsedCompanyDetails:
    values: dict[str, str | None]


@dataclass(frozen=True)
class ParsedWorkbookRow:
    row: int
    values: dict[str, Any]


@dataclass(frozen=True)
class ParsedCompanyWorkbook:
    company_details: ParsedCompanyDetails | None
    financial_years: list[ParsedWorkbookRow]
    locations: list[ParsedWorkbookRow]
    gst_registrations: list[ParsedWorkbookRow]
    gst_location_mappings: list[ParsedWorkbookRow]
    sheets: list[str]


class CompanyWorkbookError(Exception):
    def __init__(self, issues: list[CompanyImportIssue]) -> None:
        super().__init__(issues[0].message if issues else "Invalid workbook")
        self.issues = issues


def _issue(
    message: str,
    *,
    sheet: str | None = SHEET_NAME,
    row: int | None = None,
    column: str | None = None,
    field: str | None = None,
) -> CompanyImportIssue:
    return CompanyImportIssue(
        sheet=sheet,
        row=row,
        column=column,
        field=field,
        message=message,
    )


def _validate_archive(contents: bytes) -> None:
    try:
        with ZipFile(BytesIO(contents)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise CompanyWorkbookError(
                    [_issue("Workbook contains too many archive entries", sheet=None)]
                )
            total_size = 0
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise CompanyWorkbookError(
                        [_issue("Workbook contains an unsafe archive path", sheet=None)]
                    )
                if entry.flag_bits & 0x1:
                    raise CompanyWorkbookError(
                        [_issue("Encrypted workbooks are not supported", sheet=None)]
                    )
                if entry.file_size > MAX_ARCHIVE_ENTRY_BYTES:
                    raise CompanyWorkbookError(
                        [_issue("Workbook contains an oversized entry", sheet=None)]
                    )
                total_size += entry.file_size
                if total_size > MAX_ARCHIVE_BYTES:
                    raise CompanyWorkbookError(
                        [_issue("Workbook expands beyond the allowed size", sheet=None)]
                    )
                if entry.filename.casefold().endswith("vbaproject.bin"):
                    raise CompanyWorkbookError(
                        [_issue("Macro-enabled workbooks are not supported", sheet=None)]
                    )
    except BadZipFile as exc:
        raise CompanyWorkbookError(
            [_issue("File is not a valid XLSX workbook", sheet=None)]
        ) from exc


def _cell_value(value: object) -> object | None:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def _cell_text(value: object) -> str | None:
    normalized = _cell_value(value)
    return None if normalized is None else str(normalized)


def _sheet_rows(worksheet: Any) -> list[tuple[Any, ...]]:
    if worksheet.sheet_state != "visible":
        raise CompanyWorkbookError(
            [_issue(f"Worksheet '{worksheet.title}' must be visible", sheet=worksheet.title)]
        )
    if worksheet.max_row > MAX_WORKSHEET_ROWS:
        raise CompanyWorkbookError(
            [_issue("Worksheet exceeds the 100-row safety limit", sheet=worksheet.title)]
        )
    if worksheet.max_column > MAX_WORKSHEET_COLUMNS:
        raise CompanyWorkbookError(
            [_issue("Worksheet exceeds the 20-column safety limit", sheet=worksheet.title)]
        )
    rows = list(
        worksheet.iter_rows(
            min_row=1,
            max_row=max(worksheet.max_row, 2),
            max_col=max(worksheet.max_column, 1),
        )
    )
    issues: list[CompanyImportIssue] = []
    for row in rows:
        for cell in row:
            if cell.data_type == "f":
                issues.append(
                    _issue(
                        "Formulas are not allowed; provide a literal value",
                        sheet=worksheet.title,
                        row=cell.row,
                        column=cell.column_letter,
                    )
                )
            if isinstance(cell.value, str) and len(cell.value) > MAX_CELL_CHARACTERS:
                issues.append(
                    _issue(
                        "Cell value exceeds the 2048-character limit",
                        sheet=worksheet.title,
                        row=cell.row,
                        column=cell.column_letter,
                    )
                )
    if issues:
        raise CompanyWorkbookError(issues)
    return rows


def _validated_headers(
    rows: list[tuple[Any, ...]],
    *,
    sheet: str,
    columns: dict[str, str],
) -> tuple[list[str], tuple[Any, ...]]:
    header_cells = rows[0]
    headers = [_cell_text(cell.value) for cell in header_cells]
    while headers and headers[-1] is None:
        headers.pop()
    issues: list[CompanyImportIssue] = []
    seen: set[str] = set()
    for index, header in enumerate(headers, start=1):
        coordinate = header_cells[index - 1].column_letter
        if header is None:
            issues.append(
                _issue("Header cannot be blank", sheet=sheet, row=1, column=coordinate)
            )
        elif header in seen:
            issues.append(
                _issue(
                    f"Duplicate column '{header}'",
                    sheet=sheet,
                    row=1,
                    column=coordinate,
                )
            )
        elif header not in columns:
            issues.append(
                _issue(
                    f"Unsupported column '{header}'",
                    sheet=sheet,
                    row=1,
                    column=coordinate,
                )
            )
        seen.add(header or "")
    for header, field in columns.items():
        if header not in seen:
            issues.append(
                _issue(
                    f"Missing required column '{header}'",
                    sheet=sheet,
                    row=1,
                    column=header,
                    field=field,
                )
            )
    if issues:
        raise CompanyWorkbookError(issues)
    return [header for header in headers if header is not None], header_cells


def _logical_rows(
    rows: list[tuple[Any, ...]], headers: list[str], columns: dict[str, str]
) -> list[ParsedWorkbookRow]:
    parsed: list[ParsedWorkbookRow] = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = [_cell_value(row[index].value) for index in range(len(headers))]
        if not any(value is not None for value in values):
            continue
        parsed.append(
            ParsedWorkbookRow(
                row=row_number,
                values={
                    columns[header]: values[index]
                    for index, header in enumerate(headers)
                },
            )
        )
    return parsed


def _parse_company_details(worksheet: Any) -> ParsedCompanyDetails:
    rows = _sheet_rows(worksheet)
    headers, header_cells = _validated_headers(
        rows, sheet=SHEET_NAME, columns=COLUMN_TO_FIELD
    )
    logical = _logical_rows(rows, headers, COLUMN_TO_FIELD)
    if len(logical) != 1:
        row = 2 if not logical else logical[1].row
        message = (
            "Company Details must contain exactly one data row"
            if not logical
            else "Company Details must contain exactly one logical record"
        )
        raise CompanyWorkbookError([_issue(message, row=row)])
    values = {
        field: _cell_text(value) for field, value in logical[0].values.items()
    }
    if values["legal_name"] is None:
        index = headers.index("Legal Company Name")
        raise CompanyWorkbookError(
            [
                _issue(
                    "Legal Company Name is required",
                    row=logical[0].row,
                    column=header_cells[index].column_letter,
                    field="legal_name",
                )
            ]
        )
    return ParsedCompanyDetails(values=values)


def _parse_table(
    worksheet: Any,
    *,
    columns: dict[str, str],
) -> list[ParsedWorkbookRow]:
    rows = _sheet_rows(worksheet)
    headers, _ = _validated_headers(
        rows, sheet=worksheet.title, columns=columns
    )
    logical = _logical_rows(rows, headers, columns)
    if not logical:
        raise CompanyWorkbookError(
            [
                _issue(
                    "Worksheet must contain at least one data row",
                    sheet=worksheet.title,
                    row=2,
                )
            ]
        )
    return logical


def parse_company_workbook(contents: bytes) -> ParsedCompanyWorkbook:
    if not contents:
        raise CompanyWorkbookError([_issue("Uploaded workbook is empty", sheet=None)])
    if len(contents) > MAX_FILE_BYTES:
        raise CompanyWorkbookError(
            [_issue("Workbook exceeds the 5 MiB upload limit", sheet=None)]
        )
    _validate_archive(contents)
    try:
        workbook = load_workbook(
            BytesIO(contents), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise CompanyWorkbookError(
            [_issue("File is not a valid XLSX workbook", sheet=None)]
        ) from exc
    try:
        unknown = [name for name in workbook.sheetnames if name not in SUPPORTED_SHEETS]
        if unknown:
            raise CompanyWorkbookError(
                [
                    _issue(f"Unsupported worksheet '{name}'", sheet=name)
                    for name in unknown
                ]
            )
        company_details = (
            _parse_company_details(workbook[SHEET_NAME])
            if SHEET_NAME in workbook.sheetnames
            else None
        )
        financial_years = (
            _parse_table(
                workbook[FINANCIAL_YEARS_SHEET],
                columns=FINANCIAL_YEAR_COLUMN_TO_FIELD,
            )
            if FINANCIAL_YEARS_SHEET in workbook.sheetnames
            else []
        )
        locations = (
            _parse_table(workbook[LOCATIONS_SHEET], columns=LOCATION_COLUMN_TO_FIELD)
            if LOCATIONS_SHEET in workbook.sheetnames
            else []
        )
        gst_registrations = (
            _parse_table(
                workbook[GST_REGISTRATIONS_SHEET],
                columns=GST_REGISTRATION_COLUMN_TO_FIELD,
            )
            if GST_REGISTRATIONS_SHEET in workbook.sheetnames
            else []
        )
        gst_location_mappings = (
            _parse_table(
                workbook[GST_LOCATION_MAPPINGS_SHEET],
                columns=GST_LOCATION_MAPPING_COLUMN_TO_FIELD,
            )
            if GST_LOCATION_MAPPINGS_SHEET in workbook.sheetnames
            else []
        )
        return ParsedCompanyWorkbook(
            company_details=company_details,
            financial_years=financial_years,
            locations=locations,
            gst_registrations=gst_registrations,
            gst_location_mappings=gst_location_mappings,
            sheets=list(workbook.sheetnames),
        )
    finally:
        workbook.close()


def parse_company_details_workbook(contents: bytes) -> ParsedCompanyDetails:
    parsed = parse_company_workbook(contents)
    if parsed.company_details is None:
        raise CompanyWorkbookError(
            [_issue("Workbook must contain worksheet 'Company Details'", sheet=None)]
        )
    return parsed.company_details
